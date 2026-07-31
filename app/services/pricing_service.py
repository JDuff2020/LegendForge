from __future__ import annotations
from pathlib import Path
from openpyxl import load_workbook
from app.models.results import PricingSummary

class PricingFormatError(ValueError):
    pass

class PricingService:
    """Parses the user's repeated 3-column MPC layout: Card Amount / Number Decks / Cost."""

    def load(self, path: Path) -> PricingSummary:
        wb = load_workbook(path, read_only=True, data_only=True)
        try:
            for ws in wb.worksheets:
                parsed = self._parse_sheet(ws)
                if parsed.deck_sizes:
                    parsed.source_sheet = ws.title
                    return parsed
        finally:
            wb.close()
        raise PricingFormatError("No repeated Card Amount / Number Decks / Cost blocks were found.")

    def _parse_sheet(self, ws) -> PricingSummary:
        max_col = ws.max_column
        prices: dict[int, dict[str, float]] = {}
        tiers: list[str] = []
        deck_sizes: list[int] = []
        for col in range(1, max_col + 1, 3):
            headers = [str(ws.cell(1, c).value or "").strip().lower() for c in range(col, min(col + 3, max_col + 1))]
            if headers != ["card amount", "number decks", "cost"]:
                continue
            size_value = ws.cell(2, col).value
            try:
                size = int(size_value)
            except (TypeError, ValueError):
                continue
            deck_sizes.append(size)
            prices[size] = {}
            for row in range(2, ws.max_row + 1):
                tier = str(ws.cell(row, col + 1).value or "").strip()
                cost = ws.cell(row, col + 2).value
                if not tier or cost is None:
                    continue
                try:
                    prices[size][tier] = float(cost)
                except (TypeError, ValueError):
                    raise PricingFormatError(f"Invalid cost at {ws.title}!{ws.cell(row, col + 2).coordinate}")
                if tier not in tiers:
                    tiers.append(tier)
        return PricingSummary(deck_sizes=deck_sizes, quantity_tiers=tiers, prices=prices)
