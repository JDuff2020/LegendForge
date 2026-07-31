from __future__ import annotations
from collections import Counter
from pathlib import Path
from openpyxl import load_workbook
from app.models.results import InventorySummary

QUANTITY_NAMES = {"quantity", "qty", "count", "copies", "number", "cards"}
CATEGORY_NAMES = {"category", "type", "card type", "card_type", "group"}

class InventoryService:
    """Generic Phase-1 workbook profiler. Phase 3 will add Legendary-specific schemas."""

    def load(self, path: Path) -> InventorySummary:
        wb = load_workbook(path, read_only=True, data_only=True)
        sheets, total_rows, total_cards = [], 0, 0
        categories = Counter()
        detected = {}
        try:
            for ws in wb.worksheets:
                sheets.append(ws.title)
                rows = list(ws.iter_rows(values_only=True))
                nonempty = [r for r in rows if any(v not in (None, "") for v in r)]
                if not nonempty:
                    continue
                header_index = self._find_header_row(nonempty[:20])
                header = [str(v or "").strip().lower() for v in nonempty[header_index]]
                qty_col = next((i for i, h in enumerate(header) if h in QUANTITY_NAMES), None)
                cat_col = next((i for i, h in enumerate(header) if h in CATEGORY_NAMES), None)
                detected[ws.title] = {"header_row": header_index + 1, "quantity_column": qty_col, "category_column": cat_col}
                data = nonempty[header_index + 1:]
                total_rows += len(data)
                for row in data:
                    qty = 1
                    if qty_col is not None and qty_col < len(row):
                        try:
                            qty = max(0, int(float(row[qty_col] or 0)))
                        except (TypeError, ValueError):
                            qty = 0
                    total_cards += qty
                    if cat_col is not None and cat_col < len(row) and row[cat_col] not in (None, ""):
                        categories[str(row[cat_col]).strip()] += qty
        finally:
            wb.close()
        return InventorySummary(sheets=sheets, total_rows=total_rows, estimated_cards=total_cards, categories=dict(categories), detected_columns=detected)

    @staticmethod
    def _find_header_row(rows) -> int:
        best_index, best_score = 0, -1
        keywords = QUANTITY_NAMES | CATEGORY_NAMES | {"name", "card", "card name", "expansion", "set"}
        for i, row in enumerate(rows):
            score = sum(1 for v in row if str(v or "").strip().lower() in keywords)
            if score > best_score:
                best_index, best_score = i, score
        return best_index
